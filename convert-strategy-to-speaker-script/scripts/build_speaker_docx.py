#!/usr/bin/env python3
"""Build a polished consultative-sales speaker-script DOCX from structured Markdown."""

from argparse import ArgumentParser
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FONT = "Microsoft YaHei"
NAVY = "17324D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
ORANGE = "E05A33"
INK = "202A33"
MUTED = "66727E"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
WIDTH_DXA = 9360
INDENT_DXA = 120


def font(run, size=11, bold=False, italic=False, color=INK):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    for key in ("ascii", "hAnsi", "eastAsia"):
        rpr.rFonts.set(qn(f"w:{key}"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    node = pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        pr.append(node)
    node.set(qn("w:fill"), fill)


def cell_margins(cell, top=90, bottom=90, start=120, end=120):
    pr = cell._tc.get_or_add_tcPr()
    mar = pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        pr.append(mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_geometry(table, widths):
    table.autofit = False
    pr = table._tbl.tblPr
    tw = pr.find(qn("w:tblW")) or OxmlElement("w:tblW")
    if tw.getparent() is None:
        pr.append(tw)
    tw.set(qn("w:w"), str(sum(widths)))
    tw.set(qn("w:type"), "dxa")
    ti = pr.find(qn("w:tblInd")) or OxmlElement("w:tblInd")
    if ti.getparent() is None:
        pr.append(ti)
    ti.set(qn("w:w"), str(INDENT_DXA))
    ti.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            pr = cell._tc.get_or_add_tcPr()
            cw = pr.find(qn("w:tcW")) or OxmlElement("w:tcW")
            if cw.getparent() is None:
                pr.append(cw)
            cw.set(qn("w:w"), str(widths[idx]))
            cw.set(qn("w:type"), "dxa")
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    pr.append(node)


def border_bottom(paragraph, color=ORANGE, size="14"):
    pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "7")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    pr.append(borders)


def callout(paragraph):
    pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_BLUE)
    pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "22")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), ORANGE)
    borders.append(left)
    pr.append(borders)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "180")
    ind.set(qn("w:right"), "160")
    pr.append(ind)


def inline(paragraph, text, size=11, color=INK):
    for part in re.split(r"(\*\*.*?\*\*)", text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run(part[2:-2] if bold else part)
        font(run, size=size, bold=bold, color=color)


def abstract_numbering(doc, kind):
    root = doc.part.numbering_part.element
    used = [int(x.get(qn("w:abstractNumId"))) for x in root.findall(qn("w:abstractNum"))]
    abstract_id = max(used or [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal" if kind == "number" else "bullet")
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "%1." if kind == "number" else "•")
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    ppr.append(ind)
    level.extend((start, fmt, text, ppr))
    abstract.append(level)
    root.append(abstract)
    return abstract_id


def number_instance(doc, abstract_id):
    root = doc.part.numbering_part.element
    used = [int(x.get(qn("w:numId"))) for x in root.findall(qn("w:num"))]
    num_id = max(used or [0]) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    root.append(num)
    return num_id


def apply_num(paragraph, num_id):
    pr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    numpr.extend((level, nid))
    pr.append(numpr)


def page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = paragraph.add_run("SPEAKER SCRIPT  |  ")
    font(prefix, size=8.5, color=MUTED)
    run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run.extend((begin, instr, end))
    paragraph._p.append(run)


def styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def cover(doc, args):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(78)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(args.client.upper() if args.client else "STRATEGY PROPOSAL")
    font(r, size=10.5, bold=True, color=ORANGE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(args.title)
    font(r, size=26, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run(args.subtitle)
    font(r, size=15, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"提案方：{args.agency}  |  建议时长：{args.duration}")
    font(r, size=10.5, color=MUTED)
    border_bottom(p)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("“不是把方案讲完，而是帮助客户把问题说清。”")
    font(r, size=12, italic=True, color=DARK_BLUE)
    doc.add_page_break()


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", x) for x in cells):
            rows.append(cells)
        i += 1
    return rows, i


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    weights = [max(8, max((len(r[c]) if c < len(r) else 0) for r in rows)) for c in range(cols)]
    total = sum(weights)
    widths = [max(900, round(WIDTH_DXA * w / total)) for w in weights]
    widths[-1] += WIDTH_DXA - sum(widths)
    table = doc.add_table(rows=0, cols=cols)
    for ridx, data in enumerate(rows):
        cells = table.add_row().cells
        for cidx in range(cols):
            shade(cells[cidx], NAVY if ridx == 0 else WHITE if ridx % 2 else LIGHT_GRAY)
            p = cells[cidx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            inline(p, data[cidx] if cidx < len(data) else "", size=9.2, color=WHITE if ridx == 0 else INK)
            for run in p.runs:
                run.bold = ridx == 0
    repeat_header(table.rows[0])
    table_geometry(table, widths)


def build(args):
    text = Path(args.input).read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, attr, Inches(1))
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    styles(doc)
    hp = sec.header.paragraphs[0]
    r = hp.add_run(f"{args.client} · 现场提案口播稿" if args.client else "现场提案口播稿")
    font(r, size=8.5, bold=True, color=MUTED)
    border_bottom(hp, "C9D2DB", "6")
    page_field(sec.footer.paragraphs[0])
    cover(doc, args)

    number_abs = abstract_numbering(doc, "number")
    bullet_abs = abstract_numbering(doc, "bullet")
    current_list = None
    current_num_id = None
    major_seen = False
    i = 0
    label_headings = {"主稿", "核心 Insight", "顾问式追问", "顾问式确认", "顾问式校准问题", "主持提示", "主持动作", "转场", "尖锐观点"}
    while i < len(lines):
        line = lines[i].rstrip()
        if not line or line == "---":
            current_list = None
            i += 1
            continue
        if line.strip().startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            current_list = None
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            value = re.sub(r"\*\*(.*?)\*\*", r"\1", heading.group(2)).strip()
            if level == 1:
                if major_seen:
                    doc.add_page_break()
                major_seen = True
            p = doc.add_paragraph(style=f"Heading {level}")
            r = p.add_run(value)
            color = ORANGE if level == 3 and value in label_headings else NAVY if level == 1 else BLUE if level == 2 else DARK_BLUE
            font(r, size=16 if level == 1 else 13 if level == 2 else 12, bold=True, color=color)
            current_list = None
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            callout(p)
            r = p.add_run(line[2:].strip())
            font(r, size=11.5, bold=True, color=NAVY)
            current_list = None
            i += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        bulleted = re.match(r"^[-*]\s+(.*)$", line)
        if numbered or bulleted:
            kind = "number" if numbered else "bullet"
            if current_list != kind:
                current_num_id = number_instance(doc, number_abs if kind == "number" else bullet_abs)
                current_list = kind
            p = doc.add_paragraph()
            apply_num(p, current_num_id)
            inline(p, (numbered or bulleted).group(1))
            i += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.widow_control = True
        inline(p, line)
        current_list = None
        i += 1

    props = doc.core_properties
    props.title = args.title
    props.subject = args.subtitle
    props.author = args.agency
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields")) or OxmlElement("w:updateFields")
    if update.getparent() is None:
        settings.append(update)
    update.set(qn("w:val"), "true")
    doc.save(args.output)


def main():
    parser = ArgumentParser()
    parser.add_argument("--input", required=True, help="Structured UTF-8 Markdown body")
    parser.add_argument("--output", required=True, help="Output .docx path")
    parser.add_argument("--title", required=True)
    parser.add_argument("--subtitle", default="现场提案口播稿 · 顾问式销售版")
    parser.add_argument("--client", default="")
    parser.add_argument("--agency", default="提案团队")
    parser.add_argument("--duration", default="80–100 分钟")
    args = parser.parse_args()
    build(args)


if __name__ == "__main__":
    main()
